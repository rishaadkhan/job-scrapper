"""Unit tests for scoring, resume matching, Anthropic bullet generator, and Excel exporter"""
import unittest
import asyncio
import os
import tempfile
from unittest.mock import AsyncMock, patch, MagicMock
import openpyxl
from scoring.taxonomy import BACKEND_TAXONOMY, CANONICAL_SKILL_MAP
from scoring.extractor import KeywordExtractor
from scoring.matcher import ResumeMatcher
from scoring.bullet_generator import BulletGenerator
from scoring import JobScorer
from exporter import ExcelExporter


class TestTaxonomyAndKeywordExtractor(unittest.TestCase):
    """Tests for backend taxonomy and keyword extractor."""

    def setUp(self):
        self.extractor = KeywordExtractor()

    def test_canonical_mapping(self):
        self.assertEqual(CANONICAL_SKILL_MAP["k8s"], "Kubernetes")
        self.assertEqual(CANONICAL_SKILL_MAP["postgres"], "PostgreSQL")
        self.assertEqual(CANONICAL_SKILL_MAP["golang"], "Go")
        self.assertEqual(CANONICAL_SKILL_MAP["spring boot"], "Spring Boot")

    def test_extract_keywords_from_jd(self):
        title = "Senior Backend Engineer - Payments"
        jd = """
        We are looking for a Java Backend Engineer with extensive experience in Spring Boot,
        Microservices, and PostgreSQL. You will build high-throughput distributed systems
        on AWS using Docker, Kubernetes, and Apache Kafka.
        Requirements:
        - 2+ years of Java and Spring Boot.
        - Experience with Redis caching and REST APIs.
        """
        keywords = self.extractor.extract_keywords(title, jd, top_n=10)
        self.assertIn("Java", keywords)
        self.assertIn("Spring Boot", keywords)
        self.assertIn("Microservices", keywords)
        self.assertIn("PostgreSQL", keywords)
        self.assertIn("Docker", keywords)
        self.assertIn("Kafka", keywords)


class TestResumeMatcher(unittest.TestCase):
    """Tests for ResumeMatcher skill parsing and 0-100 scoring."""

    def test_resume_matching_and_scoring(self):
        sample_resume = """
        # Software Engineer
        Skills: Java, Python, Spring Boot, PostgreSQL, AWS, Docker, Microservices, REST API, Redis.
        """
        with tempfile.NamedTemporaryFile(mode="w", suffix=".md", delete=False) as tf:
            tf.write(sample_resume)
            temp_resume_path = tf.name

        try:
            matcher = ResumeMatcher(resume_path=temp_resume_path)
            self.assertIn("Java", matcher.resume_keywords)
            self.assertIn("PostgreSQL", matcher.resume_keywords)
            self.assertIn("AWS", matcher.resume_keywords)

            # High match JD
            title1 = "Java Backend Developer"
            jd1 = "Need experience in Java, Spring Boot, Microservices, PostgreSQL, and AWS."
            res1 = matcher.score_job(title1, jd1)
            self.assertGreaterEqual(res1['match_score'], 80)
            self.assertEqual(len(res1['missing_from_resume']), 0)

            # JD with missing skills (e.g. Golang, Kubernetes, Kafka, Snowflake)
            title2 = "Go Infrastructure Engineer"
            jd2 = "We build data pipelines with Go, Kubernetes, Kafka, Snowflake, and ClickHouse."
            res2 = matcher.score_job(title2, jd2)
            self.assertIn("Go", res2['missing_from_resume'])
            self.assertIn("Kafka", res2['missing_from_resume'])
            self.assertLess(res2['match_score'], 50)

        finally:
            if os.path.exists(temp_resume_path):
                os.remove(temp_resume_path)


class TestBulletGenerator(unittest.IsolatedAsyncioTestCase):
    """Tests for Anthropic Claude bullet point generation and graceful fallback."""

    async def test_skips_when_no_api_key(self):
        generator = BulletGenerator(api_key="")
        bullets = await generator.generate_bullets(
            company="Stripe",
            title="Backend Engineer",
            description="Build payment APIs in Java",
            resume_text="Java developer with 1 year exp",
            top_keywords=["Java", "Spring Boot"],
            missing_keywords=["Kafka"],
            match_score=85
        )
        self.assertEqual(bullets, "")

    async def test_skips_when_score_below_threshold(self):
        generator = BulletGenerator(api_key="sk-ant-mock-key", threshold=70)
        bullets = await generator.generate_bullets(
            company="Stripe",
            title="Frontend Engineer",
            description="React frontend",
            resume_text="Java backend",
            top_keywords=["React"],
            missing_keywords=["React"],
            match_score=30
        )
        self.assertEqual(bullets, "")

    async def test_generates_bullets_on_mock_api_success(self):
        generator = BulletGenerator(api_key="sk-ant-mock-key", threshold=60)
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.json.return_value = {
            "content": [
                {
                    "type": "text",
                    "text": "• Architected high-throughput REST microservices in Java and Spring Boot, reducing API latency by 30%.\n• Integrated Apache Kafka event streams and PostgreSQL database."
                }
            ]
        }

        mock_http_client = AsyncMock()
        mock_http_client.post.return_value = mock_resp

        bullets = await generator.generate_bullets(
            company="Stripe",
            title="Backend Engineer",
            description="Build scalable microservices",
            resume_text="Java developer with Spring Boot experience",
            top_keywords=["Java", "Spring Boot", "Kafka"],
            missing_keywords=["Kafka"],
            match_score=80,
            client=mock_http_client
        )

        self.assertIn("Architected high-throughput REST microservices", bullets)
        self.assertIn("Apache Kafka", bullets)


class TestExcelExporterSchemaAndFormatting(unittest.TestCase):
    """Tests for ExcelExporter 15-column schema, frozen panes, hyperlinks, and conditional formatting."""

    def test_export_jobs_formatting(self):
        exporter = ExcelExporter()
        sample_jobs = [
            {
                'company': 'Stripe',
                'company_type': 'Tier-1 GCC',
                'title': 'Software Engineer - Payments Backend',
                'match_score': 85,
                'experience_range': '0-2 years',
                'location': 'Bengaluru, India',
                'top_jd_keywords': ['Java', 'Spring Boot', 'PostgreSQL', 'AWS'],
                'missing_from_resume': [],
                'suggested_bullets': '• Developed Spring Boot services supporting 50k requests/sec.',
                'job_id': '123456',
                'posted_date': '2026-09-12',
                'link': 'https://stripe.com/jobs/123456',
                'portal_url': 'https://stripe.com/jobs',
                'description': 'Core payment processing backend team in Bangalore.',
                'scraped_at': '2026-09-12 10:00:00'
            },
            {
                'company': 'Thought Machine',
                'company_type': 'High-Paying Startup',
                'title': 'Platform Engineer',
                'match_score': 35,
                'experience_range': '0-3 years',
                'location': 'Bengaluru, India',
                'top_jd_keywords': ['Go', 'Kubernetes', 'gRPC'],
                'missing_from_resume': ['Go', 'Kubernetes'],
                'suggested_bullets': '',
                'job_id': '999',
                'posted_date': '2026-09-11',
                'link': 'https://thoughtmachine.net/jobs/999',
                'portal_url': 'https://thoughtmachine.net/careers',
                'description': 'Platform infrastructure team.',
                'scraped_at': '2026-09-12 10:00:00'
            }
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            with patch('exporter.OUTPUT_DIR', temp_dir):
                filepath = exporter.export_jobs(sample_jobs)
                self.assertIsNotNone(filepath)
                self.assertTrue(os.path.exists(filepath))

                # Load workbook and verify schema
                wb = openpyxl.load_workbook(filepath)
                ws = wb.active
                self.assertEqual(ws.title, "Job Leads")

                # Verify header columns (15 columns)
                headers = [cell.value for cell in ws[1]]
                self.assertEqual(len(headers), 15)
                self.assertEqual(headers[3], "Match Score")
                self.assertEqual(headers[6], "Top JD Keywords")
                self.assertEqual(headers[7], "Missing From Resume")
                self.assertEqual(headers[8], "Suggested Bullet Edits")

                # Verify frozen panes
                self.assertEqual(ws.freeze_panes, "A2")

                # Verify row data values
                self.assertEqual(ws.cell(row=2, column=1).value, "Stripe")
                self.assertEqual(ws.cell(row=2, column=4).value, 85)
                self.assertEqual(ws.cell(row=2, column=7).value, "Java, Spring Boot, PostgreSQL, AWS")
                self.assertEqual(ws.cell(row=3, column=4).value, 35)

                # Verify hyperlinks
                cell_link = ws.cell(row=2, column=12)
                self.assertIsNotNone(cell_link.hyperlink)
                self.assertEqual(cell_link.hyperlink.target, "https://stripe.com/jobs/123456")

                # Verify conditional formatting rule added
                self.assertGreaterEqual(len(ws.conditional_formatting), 1)


if __name__ == '__main__':
    unittest.main()

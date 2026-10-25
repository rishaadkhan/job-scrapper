"""Excel export functionality with professional styling, conditional formatting, and hyperlinks"""
import os
from datetime import datetime
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
from config import OUTPUT_DIR, EXCEL_FILENAME_TEMPLATE, EXCEL_COLUMNS
from filters import JobFilter


class ExcelExporter:
    """Exports structured job leads to formatted, styled Excel spreadsheets."""

    def __init__(self):
        os.makedirs(OUTPUT_DIR, exist_ok=True)

    def export_jobs(self, jobs: List[Dict[str, Any]]) -> Optional[str]:
        if not jobs:
            print("No jobs to export")
            return None

        wb = Workbook()
        ws = wb.active
        ws.title = "Job Leads"

        # 1. Header Styling
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.append(EXCEL_COLUMNS)
        ws.row_dimensions[1].height = 28

        for col_num in range(1, len(EXCEL_COLUMNS) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        # 2. Body data rows
        link_font = Font(name="Calibri", size=10, color="0563C1", underline="single")
        default_font = Font(name="Calibri", size=10)
        score_font = Font(name="Calibri", size=11, bold=True)
        center_align = Alignment(horizontal="center", vertical="top")
        left_align = Alignment(horizontal="left", vertical="top")
        wrap_align = Alignment(horizontal="left", vertical="top", wrap_text=True)

        for row_idx, job in enumerate(jobs, start=2):
            match_score = job.get('match_score', 50)
            top_keywords = ', '.join(job.get('top_jd_keywords', []))
            missing_keywords = ', '.join(job.get('missing_from_resume', []))
            suggested_bullets = job.get('suggested_bullets', '')
            exp_range = job.get('experience_range') or JobFilter.extract_experience(job.get('title', ''), job.get('description', ''))
            apply_link = job.get('link', '')
            portal_url = job.get('portal_url', '')

            row_data = [
                job.get('company', ''),
                job.get('company_type', ''),
                job.get('title', ''),
                match_score,
                exp_range,
                job.get('location', ''),
                top_keywords,
                missing_keywords,
                suggested_bullets,
                job.get('job_id', ''),
                job.get('posted_date', ''),
                apply_link,
                portal_url,
                job.get('description', ''),
                job.get('scraped_at', datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
            ]
            ws.append(row_data)
            ws.row_dimensions[row_idx].height = 22 if not suggested_bullets else 45

            # Apply cell-level formatting
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = default_font
                cell.alignment = left_align

                # Match score centered & bold
                if col_idx == 4:
                    cell.alignment = center_align
                    cell.font = score_font

                # Center align experience and date
                elif col_idx in [5, 11, 15]:
                    cell.alignment = center_align

                # Text wrap for keywords, bullets, and descriptions
                elif col_idx in [7, 8, 9, 14]:
                    cell.alignment = wrap_align

                # Hyperlink apply link
                elif col_idx == 12 and apply_link and apply_link.startswith('http'):
                    cell.hyperlink = apply_link
                    cell.font = link_font

                # Hyperlink career portal URL
                elif col_idx == 13 and portal_url and portal_url.startswith('http'):
                    cell.hyperlink = portal_url
                    cell.font = link_font

        # 3. Conditional Formatting for Match Score (Column D)
        # Green >= 70, Yellow 40-69, Red < 40
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        green_font = Font(name="Calibri", size=11, bold=True, color="006100")
        yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        yellow_font = Font(name="Calibri", size=11, bold=True, color="9C6500")
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        red_font = Font(name="Calibri", size=11, bold=True, color="9C0006")

        max_row = max(2, len(jobs) + 1)
        score_range = f"D2:D{max_row}"

        ws.conditional_formatting.add(
            score_range,
            CellIsRule(operator="greaterThanOrEqual", formula=["70"], fill=green_fill, font=green_font)
        )
        ws.conditional_formatting.add(
            score_range,
            CellIsRule(operator="between", formula=["40", "69"], fill=yellow_fill, font=yellow_font)
        )
        ws.conditional_formatting.add(
            score_range,
            CellIsRule(operator="lessThan", formula=["40"], fill=red_fill, font=red_font)
        )

        # 4. Freeze Header Row and Enable Auto-Filter
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(EXCEL_COLUMNS))}{max_row}"

        # 5. Sane Column Widths
        column_widths = {
            1: 22,   # Company Name
            2: 18,   # Company Type
            3: 32,   # Job Title
            4: 14,   # Match Score
            5: 18,   # Experience Range
            6: 25,   # Location
            7: 35,   # Top JD Keywords
            8: 30,   # Missing From Resume
            9: 45,   # Suggested Bullet Edits
            10: 20,  # Job ID
            11: 15,  # Posted Date
            12: 35,  # Official Apply Link
            13: 30,  # Career Portal URL
            14: 60,  # Full Job Description
            15: 22   # Scraped Timestamp
        }

        for col_idx, width in column_widths.items():
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width

        filename = EXCEL_FILENAME_TEMPLATE.format(date=datetime.now().strftime('%Y-%m-%d'))
        filepath = os.path.join(OUTPUT_DIR, filename)

        wb.save(filepath)
        print(f"Exported {len(jobs)} scored leads to {filepath}")

        return filepath

    def export_scrape_runs(self, runs: List[Dict[str, Any]], filename_prefix: str = "scrape_runs_report") -> str:
        """Export scrape run telemetry metrics to a formatted Excel workbook."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Scrape Runs Telemetry"

        RUN_COLUMNS = [
            "Run ID",
            "Started At",
            "Completed At",
            "Status",
            "Duration (s)",
            "Total Companies",
            "Companies With Jobs",
            "Companies Zero Jobs",
            "Total Raw Jobs",
            "Total Valid Leads",
            "Avg Match Score (%)",
            "Error Details"
        ]

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.append(RUN_COLUMNS)
        ws.row_dimensions[1].height = 28

        for col_num in range(1, len(RUN_COLUMNS) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        center_align = Alignment(horizontal="center", vertical="center")
        wrap_align = Alignment(horizontal="left", vertical="top", wrap_text=True)

        for row_idx, r in enumerate(runs, start=2):
            row_data = [
                r.get("id"),
                r.get("started_at"),
                r.get("completed_at"),
                r.get("status"),
                r.get("duration_sec"),
                r.get("total_companies"),
                r.get("companies_with_jobs"),
                r.get("companies_zero_jobs"),
                r.get("total_raw_jobs"),
                r.get("total_valid_leads"),
                r.get("avg_match_score"),
                r.get("error_log") or ""
            ]
            ws.append(row_data)
            ws.row_dimensions[row_idx].height = 22

            for col_idx in range(1, len(RUN_COLUMNS) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = Font(name="Calibri", size=10)
                if col_idx in [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11]:
                    cell.alignment = center_align
                else:
                    cell.alignment = wrap_align

        ws.freeze_panes = "A2"
        max_row = max(2, len(runs) + 1)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(RUN_COLUMNS))}{max_row}"

        col_widths = {
            1: 12,   # Run ID
            2: 22,   # Started At
            3: 22,   # Completed At
            4: 15,   # Status
            5: 14,   # Duration
            6: 18,   # Total Companies
            7: 22,   # Companies With Jobs
            8: 22,   # Companies Zero Jobs
            9: 16,   # Total Raw Jobs
            10: 18,  # Total Valid Leads
            11: 20,  # Avg Match Score
            12: 50   # Error Details
        }
        for col_idx, width in col_widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f"{filename_prefix}_{timestamp}.xlsx"
        filepath = os.path.join(OUTPUT_DIR, filename)
        wb.save(filepath)
        return filepath
